# src/crews/decision_executor.py
"""决策团队执行器 - 结果处理、集体决策投票、报告生成"""

from datetime import datetime
import json
import logging
import os
import re
from typing import Any

from src.tasks.collective_decision_maker import (
    get_decision_maker,
)

logger = logging.getLogger(__name__)

# ── 操作 → 评级/分数映射 ──────────────────────────
_RATING_MAP = {
    "强烈买入": {"rating": "强烈买入", "score": 90},
    "买入": {"rating": "买入", "score": 70},
    "持有": {"rating": "持有", "score": 50},
    "卖出": {"rating": "卖出", "score": 30},
    "强烈卖出": {"rating": "强烈卖出", "score": 10},
}

# ── 评分 → 标签映射 ──────────────────────────────
_RATING_LABEL_THRESHOLDS = (
    (80, "优秀"),
    (60, "良好"),
    (40, "一般"),
)


def prepare_decision_inputs(analysis_result: dict) -> dict:
    """准备决策输入"""
    scores = analysis_result.get("collaboration_scores", {})
    recommendation = analysis_result.get("final_recommendation", {})
    summary = generate_analysis_summary(
        analysis_result.get("company", ""),
        analysis_result.get("ticker", ""),
        scores,
        recommendation,
        {},
    )
    return {
        "company": analysis_result.get("company", ""),
        "ticker": analysis_result.get("ticker", ""),
        "scores": scores,
        "analysis_recommendations": recommendation,
        "analysis_outputs": summary,
    }


def collect_decision_outputs(tasks_output: list[Any]) -> dict[str, Any]:
    """收集决策输出"""
    results = {}
    for i, output in enumerate(tasks_output):
        if hasattr(output, "raw"):
            results[f"task_{i}"] = str(output.raw)
        elif hasattr(output, "output"):
            results[f"task_{i}"] = str(output.output)
        else:
            results[f"task_{i}"] = str(output)
    return results


def extract_final_recommendation(decision_result: dict) -> dict[str, Any]:
    """提取最终建议"""
    return {
        "action": decision_result.get("consensus", decision_result.get("result", "持有")),
        "confidence": decision_result.get("confidence", 0.5),
    }


def calculate_decision_metrics(decision_result: dict) -> dict[str, Any]:
    """计算决策指标"""
    return {
        "vote_count": decision_result.get("voter_count", 0),
        "decision_type": decision_result.get("decision_type", "weighted"),
        "confidence": decision_result.get("confidence", 0.0),
    }


def get_investment_rating(action: str) -> dict:
    """获取投资评级 — 查表"""
    return _RATING_MAP.get(action, {"rating": action, "score": 50})


def generate_analysis_summary(
    company: str,
    ticker: str,
    scores: dict[str, float],
    recommendation: dict[str, Any],
    decision_metrics: dict[str, Any],
) -> str:
    """生成分析摘要"""
    lines = [
        f"# 投资分析摘要: {company} ({ticker})",
        "",
        "## 维度评分",
        "",
    ]
    for name, score in scores.items():
        if name.startswith("_"):
            continue
        if name != "overall_score":
            lines.append(f"- **{name}**: {score:.1f}/100")
    lines.extend([
        "",
        "## 综合评分",
        f"- **overall**: {scores.get('overall_score', 0):.1f}/100",
        "",
        "## 最终建议",
        f"- **操作**: {recommendation.get('action', '持有')}",
        f"- **置信度**: {recommendation.get('confidence', 0):.1%}",
        f"- **投票人数**: {decision_metrics.get('vote_count', 0)}",
        "",
    ])
    return "\n".join(lines)


def generate_investment_report(
    company: str,
    ticker: str,
    data_collection_result: dict,
    analysis_result: dict,
    decision_result: dict,
    summary: str,
) -> str:
    """生成完整投资分析报告 — 包含数据摘要、各维度分析、评分、决策推理"""
    scores = analysis_result.get("collaboration_scores", {})
    overall = scores.get("overall_score", 0)
    recommendation = decision_result.get("final_recommendation", {})
    action = recommendation.get("action", "持有")
    confidence = recommendation.get("confidence", 0)

    emoji = _RATING_MAP.get(action, {}).get("score", 50) >= 70 and "🟢" or "🔴"

    def _bar(score: float, width: int = 20) -> str:
        filled = int(score / 100 * width)
        return "█" * filled + "░" * (width - filled)

    lines = [
        f"# 投资分析报告: {company} ({ticker})",
        "",
        f"**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "---",
        "",
        f"## {emoji} 投资结论: {action}（置信度 {confidence:.0%}）",
        "",
        "---",
        "",
        "## 综合评分",
        "",
        "| 维度 | 评分 | 评级 |",
        "|------|------|------|",
    ]
    for name, score in scores.items():
        if name.startswith("_"):
            continue
        if name == "overall_score":
            continue
        level = _get_rating_label(score)
        lines.append(f"| {name} | {score:.1f} | {level} |")
    lines.append(f"| **综合** | **{overall:.1f}** | **{_get_rating_label(overall)}** |")

    data_quality = scores.get("_data_quality", "good")
    if data_quality in ("suspicious", "unreliable"):
        lines.extend([
            "",
            "> ⚠️ **数据质量警告**: 评分基于不完整数据，可能不准确，仅供参考。",
        ])

    lines.extend([
        "",
        "```",
        f"综合评分: {_bar(overall)} {overall:.1f}/100",
        "```",
        "",
        "---",
    ])

    # ── 分析 Agent 输出 ──
    agent_outputs = analysis_result.get("agent_outputs", {})
    if agent_outputs:
        _TASK_LABELS = {
            "task_0": "基本面分析",
            "task_1": "风险评估",
            "task_2": "行业分析",
            "task_3": "分析协调",
        }
        lines.append("## 各维度分析")
        lines.append("")
        for key in sorted(agent_outputs.keys()):
            label = _TASK_LABELS.get(key, key)
            text = str(agent_outputs[key]).strip()
            if text and len(text) > 20:
                lines.append(f"### {label}")
                lines.append("")
                lines.append(text)
                lines.append("")
        lines.append("---")
        lines.append("")

    # ── 决策推理 ──
    decision_outputs = decision_result.get("decision_outputs", {})
    if decision_outputs:
        lines.append("## 投资决策推理")
        lines.append("")
        for key, text in decision_outputs.items():
            text = str(text).strip()
            if text and len(text) > 20:
                lines.append(f"### 决策者 {key}")
                lines.append("")
                lines.append(text)
                lines.append("")
        lines.append("---")
        lines.append("")

    # ── 数据采集状态 ──
    lines.extend([
        "## 数据采集状态",
        "",
        f"- 数据采集: {'✅ 成功' if data_collection_result.get('status') == 'success' else '❌ 失败'}",
        f"- 数据时间: {data_collection_result.get('timestamp', 'N/A')}",
        f"- 采集耗时: {data_collection_result.get('execution_time', 0):.1f} 秒",
        "",
        "---",
        "",
        "## 决策详情",
        f"- 决策类型: {decision_result.get('collective_decision_metrics', {}).get('decision_type', 'N/A')}",
        f"- 投票人数: {decision_result.get('collective_decision_metrics', {}).get('vote_count', 0)}",
        "",
        "---",
        "",
        "## 免责声明",
        "本报告仅供参考，不构成投资建议。投资有风险，入市需谨慎。",
        "",
        "**报告由 AI 投资分析系统自动生成**",
    ])
    return "\n".join(lines)


def _get_rating_label(score: float) -> str:
    """评分映射到标签 — 阈值查表"""
    for threshold, label in _RATING_LABEL_THRESHOLDS:
        if score >= threshold:
            return label
    return "较差"


def analyze_collective_decision(decision_outputs: dict[str, str]) -> dict[str, Any]:
    """解析决策输出文本，聚合为集体决策共识

    从各 Agent 的文本输出中提取操作建议，通过 CollectiveDecisionMaker 投票聚合。
    """
    decision_maker = get_decision_maker()
    decision_maker.clear_history()

    _ACTION_PATTERN = re.compile(r"(强烈买入|买入|持有|卖出|强烈卖出)")
    _CONF_PATTERN = re.compile(r"置信度[：:]\s*(\d+(?:\.\d+)?)")
    _SCORE_PATTERN = re.compile(r"评分[：:]\s*(\d+(?:\.\d+)?)")

    votes = []
    for task_key, text in decision_outputs.items():
        action_match = _ACTION_PATTERN.search(text)
        conf_match = _CONF_PATTERN.search(text) or _SCORE_PATTERN.search(text)
        action = action_match.group(1) if action_match else "持有"

        if conf_match:
            raw_value = float(conf_match.group(1))
            confidence = raw_value / 100 if raw_value > 1 else raw_value
        else:
            confidence = 0.5
        confidence = max(0.1, min(confidence, 1.0))
        votes.append(decision_maker.cast_vote(task_key, action, confidence, text[:200]))

    if not votes:
        votes.append(decision_maker.cast_vote("default", "持有", 0.5, "默认投票"))

    return decision_maker.decide(votes)


def save_report(report: str, company: str, ticker: str) -> dict[str, str]:
    """保存报告到文件（md + 可选 docx）

    Returns:
        {"md": "path/to/report.md", "docx": "path/to/report.docx"}
    """
    os.makedirs("reports", exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"{company}_{ticker}_{timestamp}"

    md_path = f"reports/analysis_{safe_name}.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(report)

    docx_path = ""
    try:
        from docx import Document

        doc = Document()
        for line in report.split("\n"):
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.strip():
                doc.add_paragraph(line)
        docx_path = f"reports/analysis_{safe_name}.docx"
        doc.save(docx_path)
    except ImportError:
        logger.debug("python-docx 未安装，跳过 docx 生成")

    return {"md": md_path, "docx": docx_path}


def export_to_json(data: dict[str, Any], company: str, ticker: str) -> str:
    """导出数据为 JSON 文件

    Returns:
        文件路径
    """
    os.makedirs("reports", exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"{company}_{ticker}_{timestamp}"
    json_path = f"reports/data_{safe_name}.json"

    def _serialize(obj: Any) -> Any:
        if hasattr(obj, "isoformat"):
            return obj.isoformat()
        if hasattr(obj, "__dict__"):
            return str(obj)
        return str(obj)

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=_serialize)

    return json_path