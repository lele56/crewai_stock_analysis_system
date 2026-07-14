# src/tools/reporting_tools.py
"""
报告生成工具
支持多种报告格式：Markdown、JSON、HTML、Word(docx)、CSV
"""
from crewai.tools import BaseTool as CrewAIBaseTool
from typing import Dict, Any, List, Optional
import json
import os
from datetime import datetime
import logging

from src.tools.report_templates import ReportTemplates
from src.config import Config

logger = logging.getLogger(__name__)


class BaseTool(CrewAIBaseTool):
    """工具基类"""
    name: str = "Base Tool"
    description: str = "基础工具类"

    def _run(self, *args, **kwargs):
        raise NotImplementedError("子类必须实现_run方法")


class ReportWritingTool(BaseTool):
    """报告编写工具"""

    name: str = "Report Writing Tool"
    description: str = "生成标准化的投资分析报告和文档"

    def _run(self, report_data: str, report_type: str = "investment_analysis") -> str:
        try:
            logger.info(f"生成报告，类型: {report_type}")
            try:
                data = json.loads(report_data)
            except json.JSONDecodeError as e:
                logger.warning(f"JSON解析失败，尝试作为纯文本处理: {str(e)[:80]}")
                data = {"raw_content": report_data}
            if report_type == "investment_analysis":
                report = self._generate_investment_analysis_report(data)
            elif report_type == "summary":
                report = self._generate_summary_report(data)
            elif report_type == "executive_brief":
                report = self._generate_executive_brief(data)
            elif report_type == "detailed":
                report = self._generate_detailed_report(data)
            else:
                report = self._generate_generic_report(data)
            logger.info("报告生成完成")
            return report
        except Exception as e:
            error_msg = f"生成报告失败: {str(e)}"
            logger.error(error_msg)
            return error_msg

    def _generate_investment_analysis_report(self, data: Dict) -> str:
        return ReportTemplates.investment_analysis_template(data)

    def _generate_summary_report(self, data: Dict) -> str:
        return ReportTemplates.summary_template(data)

    def _generate_executive_brief(self, data: Dict) -> str:
        return ReportTemplates.executive_brief_template(data)

    def _generate_detailed_report(self, data: Dict) -> str:
        return self._generate_investment_analysis_report(data)

    def _generate_generic_report(self, data: Dict) -> str:
        return f"""# 分析报告
**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
## 报告内容
{json.dumps(data, ensure_ascii=False, indent=2)}"""


class DataExportTool(BaseTool):
    """数据导出工具 - 支持JSON/CSV/Excel/Markdown/Word/HTML/TXT"""

    name: str = "Data Export Tool"
    description: str = "将分析数据导出为多种格式"

    def _run(self, export_data: str, export_format: str = "json", filename: str = "") -> str:
        try:
            try:
                data = json.loads(export_data)
            except json.JSONDecodeError:
                data = {"raw_content": export_data}
            logger.info(f"导出数据，格式: {export_format}")
            os.makedirs('data/exports', exist_ok=True)
            if not filename:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                filename = f"export_data_{timestamp}"
            if export_format == "json":
                return self._export_json(data, filename)
            elif export_format == "csv":
                return self._export_csv(data, filename)
            elif export_format == "excel":
                return self._export_excel(data, filename)
            elif export_format == "markdown":
                return self._export_markdown(data, filename)
            elif export_format == "word":
                return self._export_word(data, filename)
            elif export_format == "html":
                return self._export_html(data, filename)
            elif export_format == "txt":
                return self._export_txt(data, filename)
            else:
                raise ValueError(f"不支持的导出格式: {export_format}")
        except Exception as e:
            error_msg = f"数据导出失败: {str(e)}"
            logger.error(error_msg)
            return error_msg

    def _export_json(self, data: Dict, filename: str) -> str:
        filepath = os.path.join('data/exports', f"{filename}.json")
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info(f"JSON导出完成: {filepath}")
        return filepath

    def _export_csv(self, data: Dict, filename: str) -> str:
        import pandas as pd
        filepath = os.path.join('data/exports', f"{filename}.csv")
        df = pd.DataFrame([data])
        df.to_csv(filepath, index=False, encoding='utf-8-sig')
        logger.info(f"CSV导出完成: {filepath}")
        return filepath

    def _export_excel(self, data: Dict, filename: str) -> str:
        import pandas as pd
        filepath = os.path.join('data/exports', f"{filename}.xlsx")
        df = pd.DataFrame([data])
        df.to_excel(filepath, index=False)
        logger.info(f"Excel导出完成: {filepath}")
        return filepath

    def _export_markdown(self, data: Dict, filename: str) -> str:
        filepath = os.path.join('data/exports', f"{filename}.md")
        md = self._dict_to_markdown(data, level=1)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(md)
        logger.info(f"Markdown导出完成: {filepath}")
        return filepath

    def _export_word(self, data: Dict, filename: str) -> str:
        filepath = os.path.join('data/exports', f"{filename}.docx")
        try:
            from docx import Document
            doc = Document()
            doc.add_heading('投资分析报告', 0)
            doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            self._dict_to_docx(data, doc)
            doc.save(filepath)
        except ImportError:
            md_content = self._dict_to_markdown(data)
            filepath = filepath.replace('.docx', '.md')
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(f"# 投资分析报告\n\n> 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n{md_content}")
            logger.info("python-docx未安装，已导出为Markdown格式")
        logger.info(f"Word导出完成: {filepath}")
        return filepath

    def _export_html(self, data: Dict, filename: str) -> str:
        filepath = os.path.join('data/exports', f"{filename}.html")
        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head><meta charset="UTF-8"><title>投资分析报告</title>
<style>body{{font-family:Arial,sans-serif;max-width:900px;margin:0 auto;padding:20px}}
h1{{color:#1a5276}}h2{{color:#2980b9;border-bottom:2px solid #2980b9}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:8px;text-align:left}}
th{{background-color:#2980b9;color:white}}</style></head>
<body><h1>投资分析报告</h1>
<p>生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
{self._dict_to_html(data)}</body></html>"""
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(html)
        logger.info(f"HTML导出完成: {filepath}")
        return filepath

    def _export_txt(self, data: Dict, filename: str) -> str:
        filepath = os.path.join('data/exports', f"{filename}.txt")
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(self._dict_to_text(data))
        logger.info(f"TXT导出完成: {filepath}")
        return filepath

    def _dict_to_markdown(self, data: Dict, level: int = 1) -> str:
        lines = []
        for key, value in data.items():
            if isinstance(value, dict):
                lines.append(f"{'#' * level} {key}")
                lines.append(self._dict_to_markdown(value, min(level + 1, 6)))
            elif isinstance(value, list):
                lines.append(f"{'#' * level} {key}")
                for item in value:
                    if isinstance(item, dict):
                        lines.append(self._dict_to_markdown(item, min(level + 1, 6)))
                    else:
                        lines.append(f"- {item}")
            else:
                lines.append(f"- **{key}**: {value}")
        return "\n".join(lines) + "\n"

    def _dict_to_html(self, data: Dict, level: int = 2) -> str:
        parts = []
        for key, value in data.items():
            if isinstance(value, dict):
                parts.append(f"<h{level}>{key}</h{level}>")
                parts.append(self._dict_to_html(value, min(level + 1, 6)))
            elif isinstance(value, list):
                parts.append(f"<h{level}>{key}</h{level}><ul>")
                for item in value:
                    parts.append(f"<li>{item}</li>")
                parts.append("</ul>")
            else:
                parts.append(f"<p><strong>{key}:</strong> {value}</p>")
        return "\n".join(parts)

    def _dict_to_text(self, data: Dict) -> str:
        return json.dumps(data, ensure_ascii=False, indent=2)

    def _dict_to_docx(self, data: Dict, doc, level: int = 1) -> None:
        from docx import Document
        for key, value in data.items():
            if isinstance(value, dict):
                doc.add_heading(key, level=min(level, 3))
                self._dict_to_docx(value, doc, level + 1)
            elif isinstance(value, list):
                doc.add_heading(key, level=min(level, 3))
                for item in value:
                    doc.add_paragraph(str(item), style='List Bullet')
            else:
                doc.add_paragraph(f"{key}: {value}")


class ReportTemplateTool(BaseTool):
    """报告模板工具"""

    name: str = "Report Template Tool"
    description: str = "使用预定义模板生成标准化报告"

    def _run(self, template_data: str, template_name: str = "standard") -> str:
        try:
            try:
                data = json.loads(template_data)
            except json.JSONDecodeError:
                data = {"raw_content": template_data}
            logger.info(f"使用模板生成报告: {template_name}")
            templates = ReportTemplates.get_templates()
            if template_name not in templates:
                raise ValueError(f"未知模板: {template_name}")
            report = templates[template_name](data)
            logger.info("模板报告生成完成")
            return report
        except Exception as e:
            error_msg = f"模板报告生成失败: {str(e)}"
            logger.error(error_msg)
            return error_msg